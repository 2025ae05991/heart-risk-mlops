import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

root = Path(__file__).resolve().parents[1]
doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.7)
section.bottom_margin = Inches(0.65)
section.left_margin = Inches(0.8)
section.right_margin = Inches(0.8)
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.05
for name, size in [
    ("Title", 27),
    ("Subtitle", 12),
    ("Heading 1", 18),
    ("Heading 2", 13),
    ("Heading 3", 11.5),
]:
    style = doc.styles[name]
    style.font.name = "Calibri"
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.bold = name != "Subtitle"
    style.paragraph_format.space_before = Pt(10 if name != "Title" else 0)
    style.paragraph_format.space_after = Pt(9)
    for el in list(style.element.xpath(".//w:color")):
        for attr in ["themeColor", "themeTint", "themeShade"]:
            el.attrib.pop(qn("w:" + attr), None)
    ppr = style.element.find(qn("w:pPr"))
    if ppr is not None:
        border = ppr.find(qn("w:pBdr"))
        if border is not None:
            ppr.remove(border)

footer = section.footer.paragraphs[0]
footer.paragraph_format.space_before = Pt(0)
footer.paragraph_format.space_after = Pt(0)
run = footer.add_run("Heart Risk MLOps   Docker Desktop Setup and Run Guide")
run.font.size = Pt(8)
run.font.color.rgb = RGBColor.from_string("555555")
footer.add_run(" " * 6 + "Page ").font.size = Pt(8)
fld = OxmlElement("w:fldSimple")
fld.set(qn("w:instr"), "PAGE")
footer._p.append(fld)


def rich(paragraph, text):
    for piece in re.split(r"(`[^`]+`)", text):
        if piece.startswith("`") and piece.endswith("`"):
            run = paragraph.add_run(piece[1:-1])
            run.font.name = "Courier New"
            run.font.size = Pt(9.5)
        else:
            paragraph.add_run(piece)


def add_table(lines):
    rows = [[c.strip() for c in line.strip().strip("|").split("|")] for line in lines]
    rows = [r for r in rows if not all(re.fullmatch(r"[:\- ]+", x or "-") for x in r)]
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.autofit = False
    widths = [2.15, 4.75] if len(rows[0]) == 2 else [6.9 / len(rows[0])] * len(rows[0])
    for row_idx, values in enumerate(rows):
        row = table.rows[row_idx]
        pr = row._tr.get_or_add_trPr()
        nosplit = OxmlElement("w:cantSplit")
        pr.append(nosplit)
        if row_idx == 0:
            repeat = OxmlElement("w:tblHeader")
            pr.append(repeat)
        for col_idx, value in enumerate(values):
            cell = row.cells[col_idx]
            cell.width = Inches(widths[col_idx])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            borders = OxmlElement("w:tcBorders")
            for edge in ["top", "left", "bottom", "right"]:
                el = OxmlElement("w:" + edge)
                el.set(qn("w:val"), "single")
                el.set(qn("w:sz"), "4")
                el.set(qn("w:color"), "D9D9D9")
                borders.append(el)
            tcpr.append(borders)
            margins = OxmlElement("w:tcMar")
            for edge in ["top", "left", "bottom", "right"]:
                el = OxmlElement("w:" + edge)
                el.set(qn("w:w"), "90")
                el.set(qn("w:type"), "dxa")
                margins.append(el)
            tcpr.append(margins)
            shade = OxmlElement("w:shd")
            shade.set(
                qn("w:fill"),
                "203C50" if row_idx == 0 else ("F1F4F6" if row_idx % 2 == 0 else "FFFFFF"),
            )
            tcpr.append(shade)
            para = cell.paragraphs[0]
            para.paragraph_format.space_after = Pt(1)
            para.paragraph_format.line_spacing = 1.0
            rich(para, value)
            for run in para.runs:
                run.font.size = Pt(10)
                if row_idx == 0:
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


lines = (root / "docs/Setup_and_Run_Guide.md").read_text().splitlines()
i = 0
while i < len(lines):
    line = lines[i]
    if not line.strip():
        i += 1
        continue
    if line == "<!-- pagebreak -->":
        doc.add_page_break()
        i += 1
        continue
    if line.startswith("# "):
        rich(doc.add_paragraph(style="Title"), line[2:])
        i += 1
        continue
    if line.startswith("## "):
        rich(doc.add_paragraph(style="Heading 1"), line[3:])
        i += 1
        continue
    if line.startswith("### "):
        rich(doc.add_paragraph(style="Heading 2"), line[4:])
        i += 1
        continue
    if line.startswith("```"):
        i += 1
        while i < len(lines) and not lines[i].startswith("```"):
            para = doc.add_paragraph()
            para.paragraph_format.space_after = Pt(0)
            para.paragraph_format.line_spacing = 1.0
            run = para.add_run(lines[i])
            run.font.name = "Courier New"
            run.font.size = Pt(9.4)
            i += 1
        doc.add_paragraph().paragraph_format.space_after = Pt(0)
        i += 1
        continue
    if line.startswith("|"):
        batch = []
        while i < len(lines) and lines[i].startswith("|"):
            batch.append(lines[i])
            i += 1
        add_table(batch)
        continue
    if re.match(r"^\d+\. ", line):
        text = line
        i += 1
        while i < len(lines) and lines[i].startswith("   "):
            text += " " + lines[i].strip()
            i += 1
        para = doc.add_paragraph()
        para.paragraph_format.left_indent = Inches(0.16)
        para.paragraph_format.first_line_indent = Inches(-0.16)
        rich(para, text)
        continue
    text = line
    i += 1
    while (
        i < len(lines)
        and lines[i].strip()
        and not lines[i].startswith(("#", "|", "```", "<!--"))
        and not re.match(r"^\d+\. ", lines[i])
    ):
        text += " " + lines[i].strip()
        i += 1
    rich(doc.add_paragraph(), text)

doc.core_properties.title = "Heart Risk MLOps Setup and Run Guide"
doc.core_properties.author = "Ameer Basha"
doc.core_properties.subject = "Windows PowerShell VS Code and Docker Desktop Kubernetes"
path = root / "docs/Docker_Desktop_Setup_and_Run_Guide.docx"
doc.save(path)
print(path)
